import pytest
import os
import shutil
from sass_legal.core.document_converter import convert_to_markdown

# Define the directory for test files and ensure it exists
TEST_FILE_DIR = os.path.join(os.path.dirname(__file__), "test_files")
# Dummy file names
DUMMY_TXT = os.path.join(TEST_FILE_DIR, "dummy.txt")
DUMMY_DOCX = os.path.join(TEST_FILE_DIR, "dummy.docx")
DUMMY_PDF = os.path.join(TEST_FILE_DIR, "dummy.pdf")
UNSUPPORTED_FILE = os.path.join(TEST_FILE_DIR, "dummy.unsupported")

@pytest.fixture(scope="module", autouse=True)
def setup_test_files():
    """Create dummy files for testing."""
    os.makedirs(TEST_FILE_DIR, exist_ok=True)

    # Create dummy TXT (already created by previous step, but good to have here for clarity)
    with open(DUMMY_TXT, "w") as f:
        f.write("This is a dummy text file.\nIt contains multiple lines.\nFor testing purposes.")

    # Create dummy DOCX (already created, ensure content for consistency if needed)
    # from docx import Document
    # doc = Document()
    # doc.add_paragraph("This is a dummy DOCX file.")
    # doc.add_paragraph("Created for testing document conversion.")
    # doc.save(DUMMY_DOCX)

    # Create dummy PDF (already created, ensure content for consistency if needed)
    # from reportlab.pdfgen import canvas
    # from reportlab.lib.pagesizes import letter
    # c = canvas.Canvas(DUMMY_PDF, pagesize=letter)
    # c.drawString(100, 750, "This is a dummy PDF file.")
    # c.drawString(100, 730, "Generated for testing purposes.")
    # c.save()

    # Create a dummy unsupported file
    with open(UNSUPPORTED_FILE, "w") as f:
        f.write("This is an unsupported file type.")

    yield
    # Teardown: Remove the test_files directory after tests are done
    # Comment out for inspection if tests fail. If running tests in parallel or frequently,
    # it might be better to let a global test setup/teardown handle this directory.
    # For now, this module-scoped fixture handles its own files.
    # if os.path.exists(TEST_FILE_DIR): # Ensure cleanup only if directory exists
    #     shutil.rmtree(TEST_FILE_DIR)


def test_convert_txt_to_markdown():
    """Test conversion of a valid TXT file."""
    assert os.path.exists(DUMMY_TXT), "Dummy TXT file should exist"
    markdown_content = convert_to_markdown(DUMMY_TXT)
    assert isinstance(markdown_content, str)
    assert len(markdown_content) > 0
    assert "dummy text file" in markdown_content.lower()

def test_convert_docx_to_markdown():
    """Test conversion of a valid DOCX file."""
    assert os.path.exists(DUMMY_DOCX), "Dummy DOCX file should exist"
    markdown_content = convert_to_markdown(DUMMY_DOCX)
    assert isinstance(markdown_content, str)
    assert len(markdown_content) > 0
    assert "dummy DOCX file" in markdown_content

def test_convert_pdf_to_markdown():
    """Test conversion of a valid PDF file."""
    assert os.path.exists(DUMMY_PDF), "Dummy PDF file should exist"
    markdown_content = convert_to_markdown(DUMMY_PDF)
    assert isinstance(markdown_content, str)
    assert len(markdown_content) > 0
    assert "dummy PDF file" in markdown_content

# --- Tests for error handling and edge cases ---

def test_convert_file_not_found():
    """Test conversion with a non-existent file, expecting FileNotFoundError."""
    with pytest.raises(FileNotFoundError):
        convert_to_markdown("non_existent_file.txt")

def test_convert_unsupported_format():
    """Test conversion with an unsupported file format, expecting UnsupportedFileTypeError."""
    assert os.path.exists(UNSUPPORTED_FILE), "Dummy unsupported file should exist"
    # Assuming convert_to_markdown now raises UnsupportedFileTypeError from .errors
    from sass_legal.core.errors import UnsupportedFileTypeError
    with pytest.raises(UnsupportedFileTypeError) as excinfo:
        convert_to_markdown(UNSUPPORTED_FILE)
    assert "Unsupported file format: .unsupported" in str(excinfo.value)

# --- Tests for empty files ---

def test_convert_empty_txt_file():
    """Test conversion of an empty TXT file."""
    empty_txt_path = os.path.join(TEST_FILE_DIR, "empty.txt")
    with open(empty_txt_path, "w") as f:
        f.write("")
    markdown_content = convert_to_markdown(empty_txt_path)
    assert isinstance(markdown_content, str)
    assert markdown_content == "" # Empty TXT should result in empty markdown

def test_convert_empty_docx_file():
    """Test conversion of an empty DOCX file."""
    # Creating a truly empty (but valid) DOCX is tricky without special tools or libraries.
    # A DOCX file created by `docx.Document().save()` is not truly empty but has minimal structure.
    from docx import Document
    empty_docx_path = os.path.join(TEST_FILE_DIR, "empty.docx")
    doc = Document()
    doc.save(empty_docx_path)

    markdown_content = convert_to_markdown(empty_docx_path)
    assert isinstance(markdown_content, str)
    assert markdown_content.strip() == "" # Expecting no visible text content

def test_convert_empty_pdf_file():
    """Test conversion of an empty PDF file."""
    # Similar to DOCX, a truly "empty" PDF that is still valid needs minimal structure.
    # PyPDF2 might extract nothing or raise an error depending on the PDF.
    # Let's create a PDF with no actual text content.
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import letter
    empty_pdf_path = os.path.join(TEST_FILE_DIR, "empty.pdf")
    c = canvas.Canvas(empty_pdf_path, pagesize=letter)
    c.save() # Creates a PDF with one blank page

    markdown_content = convert_to_markdown(empty_pdf_path)
    assert isinstance(markdown_content, str)
    # PyPDF2's extract_text() on a blank page often results in empty string or just newlines
    assert markdown_content.strip() == ""

# --- Tests for potentially corrupted/malformed files ---
# These tests are conceptual and depend on how PyPDF2/python-docx handle errors.
# Creating truly "corrupted" files that trigger specific library errors is hard.
# We'll simulate by creating files that are text but have PDF/DOCX extensions.

def test_convert_corrupted_pdf_as_txt():
    """Test conversion of a TXT file named as PDF, expecting DocumentProcessingError."""
    corrupted_pdf_path = os.path.join(TEST_FILE_DIR, "corrupted.pdf")
    with open(corrupted_pdf_path, "w") as f:
        f.write("This is not a real PDF, just plain text.")

    from sass_legal.core.errors import DocumentProcessingError
    with pytest.raises(DocumentProcessingError):
        # PyPDF2 will likely fail to parse this as a PDF.
        convert_to_markdown(corrupted_pdf_path)

def test_convert_corrupted_docx_as_txt():
    """Test conversion of a TXT file named as DOCX, expecting DocumentProcessingError."""
    corrupted_docx_path = os.path.join(TEST_FILE_DIR, "corrupted.docx")
    with open(corrupted_docx_path, "w") as f:
        f.write("This is not a real DOCX, just plain text.")

    from sass_legal.core.errors import DocumentProcessingError
    with pytest.raises(DocumentProcessingError):
        # python-docx will likely fail to parse this as a DOCX.
        convert_to_markdown(corrupted_docx_path)


if __name__ == "__main__":
    pytest.main([__file__])
